"""Evaluate TCEP concentrations for robust DCPIP assay Z' controls.

Expected 384-well endpoint layout (600 nm):
    * D1-D12 -> pH 6
    * E1-E12 -> pH 7
    * F1-F12 -> pH 8
    * cols 1-3 -> 0 mM TCEP / water (negative controls)
    * cols 4-6 -> 1 mM TCEP (positive controls)
    * cols 7-9 -> 5 mM TCEP (positive controls)
    * cols 10-12 -> 10 mM TCEP (positive controls)

For each pH and non-zero TCEP concentration, the script calculates:
    Z' = 1 - 3 * (sigma_positive + sigma_negative) / |mu_positive - mu_negative|

Usage
-----
    python tcep_zprime.py test/260824_dcpip_tcep_test_0_1_5_10mM.xlsx
"""

from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import openpyxl


DEFAULT_OUTPUT_SUBDIR = "outputs"
WAVELENGTH_NM = 600
PH_ROWS: dict[str, float] = {"D": 6.0, "E": 7.0, "F": 8.0}
TCEP_COLUMNS: dict[float, tuple[int, int]] = {
    0.0: (1, 3),
    1.0: (4, 6),
    5.0: (7, 9),
    10.0: (10, 12),
}

NATURE_RC = {
    "svg.fonttype": "none",
    "pdf.fonttype": 42,
    "font.family": "sans-serif",
    "font.sans-serif": ["Helvetica", "Arial", "DejaVu Sans"],
    "font.size": 6,
    "axes.labelsize": 6,
    "axes.titlesize": 6,
    "xtick.labelsize": 6,
    "ytick.labelsize": 6,
    "legend.fontsize": 6,
    "axes.linewidth": 0.5,
    "xtick.major.width": 0.5,
    "ytick.major.width": 0.5,
    "xtick.major.size": 2.0,
    "ytick.major.size": 2.0,
    "axes.spines.top": False,
    "axes.spines.right": False,
}
NATURE_PANEL_INCHES = (7.08, 5.20)
NATURE_WIDE_PANEL_INCHES = (7.08, 2.35)
NATURE_COMPOSITE_PANEL_INCHES = (7.08, 3.70)

WATER_BG = "#dce9f7"
TCEP_BG = "#d9ecd0"
PH_COLORS = {6.0: "#1f77b4", 7.0: "#2ca02c", 8.0: "#d62728"}
TCEP_COLORS = {0.0: "#1f77b4", 1.0: "#7a9cc6", 5.0: "#5ca66b", 10.0: "#2f7d44"}


@dataclass(frozen=True)
class ZPrimeResult:
    ph: float
    tcep_mM: float
    negative_values: np.ndarray
    positive_values: np.ndarray
    negative_mean: float
    negative_std: float
    positive_mean: float
    positive_std: float
    z_prime: float


def _find_plate_grid(ws) -> tuple[int, int]:
    """Return the header row and zero-based first data-column index for a 1..24 grid."""
    for row_idx, values in enumerate(ws.iter_rows(values_only=True), start=1):
        for start in range(len(values) - 1):
            if values[start] == 1 and values[start + 1] == 2:
                return row_idx, start
    raise ValueError("Could not find the plate column-header row (1, 2, ...).")


def load_endpoint_plate(xlsx_path: Path, sheet: str | None = None) -> dict[str, np.ndarray]:
    """Load a BioTek endpoint plate export as ``{row: 24 absorbance values}``."""
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    worksheet = workbook[sheet] if sheet else workbook[workbook.sheetnames[0]]
    header_row, first_data_col = _find_plate_grid(worksheet)

    plate: dict[str, np.ndarray] = {}
    # The export can contain more than one 1..24 grid (e.g. an endpoint read
    # followed by an absorbance spectrum). Restrict parsing to this grid's 16
    # plate rows so later results cannot overwrite the endpoint values.
    for row in worksheet.iter_rows(min_row=header_row + 1, max_row=header_row + 16,
                                  values_only=True):
        row_label = row[first_data_col - 1]
        if not isinstance(row_label, str) or len(row_label) != 1 or not row_label.isalpha():
            continue
        values = [float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else np.nan
                  for v in row[first_data_col:first_data_col + 24]]
        plate[row_label.upper()] = np.asarray(values, dtype=float)
    return plate


def calculate_z_prime(negative: np.ndarray, positive: np.ndarray) -> tuple[float, float, float, float, float]:
    """Return negative mean/std, positive mean/std, and standard Z' factor."""
    negative = negative[np.isfinite(negative)]
    positive = positive[np.isfinite(positive)]
    if negative.size < 2 or positive.size < 2:
        return np.nan, np.nan, np.nan, np.nan, np.nan

    negative_mean = float(np.mean(negative))
    positive_mean = float(np.mean(positive))
    negative_std = float(np.std(negative, ddof=1))
    positive_std = float(np.std(positive, ddof=1))
    separation = abs(positive_mean - negative_mean)
    z_prime = (1 - 3 * (positive_std + negative_std) / separation
               if separation > 0 else np.nan)
    return negative_mean, negative_std, positive_mean, positive_std, float(z_prime)


def analyze_tcep_z_prime(plate: dict[str, np.ndarray]) -> list[ZPrimeResult]:
    """Calculate Z' for each non-zero TCEP concentration at each tested pH."""
    neg_start, neg_end = TCEP_COLUMNS[0.0]
    results: list[ZPrimeResult] = []

    for row, ph in PH_ROWS.items():
        if row not in plate:
            raise ValueError(f"Expected pH {ph:g} data in row {row}, but that row was not found.")
        negative = plate[row][neg_start - 1:neg_end]
        for tcep_mM, (start_col, end_col) in TCEP_COLUMNS.items():
            if tcep_mM == 0.0:
                continue
            positive = plate[row][start_col - 1:end_col]
            neg_mean, neg_std, pos_mean, pos_std, z_prime = calculate_z_prime(negative, positive)
            results.append(ZPrimeResult(
                ph=ph, tcep_mM=tcep_mM,
                negative_values=negative, positive_values=positive,
                negative_mean=neg_mean, negative_std=neg_std,
                positive_mean=pos_mean, positive_std=pos_std,
                z_prime=z_prime,
            ))
    return results


def plot_tcep_z_prime(plate: dict[str, np.ndarray], results: list[ZPrimeResult], out_path: Path,
                       title: str = "") -> None:
    """Plot raw controls by pH above a compact Z' comparison panel."""
    plt.rcParams.update(NATURE_RC)
    figure = plt.figure(figsize=NATURE_COMPOSITE_PANEL_INCHES)
    grid = figure.add_gridspec(2, 3, height_ratios=(1.25, 0.75), hspace=0.23, wspace=0.30)
    concentration_labels = ["0", "1", "5", "10"]

    for axis_index, (row, ph) in enumerate(PH_ROWS.items()):
        ax = figure.add_subplot(grid[0, axis_index])
        values = plate[row]
        group_means, group_stds = [], []
        for position, (tcep_mM, (start_col, end_col)) in enumerate(TCEP_COLUMNS.items()):
            group = values[start_col - 1:end_col]
            group = group[np.isfinite(group)]
            if tcep_mM == 0:
                ax.axvspan(position - 0.45, position + 0.45, color=WATER_BG, zorder=0)
            else:
                ax.axvspan(position - 0.45, position + 0.45, color=TCEP_BG, alpha=0.7, zorder=0)
            jitter = np.linspace(-0.10, 0.10, len(group)) if len(group) > 1 else np.zeros(len(group))
            ax.plot(np.full(len(group), position) + jitter, group, "o", ms=3,
                    color=TCEP_COLORS[tcep_mM], mec="none", zorder=3)
            group_means.append(np.mean(group) if group.size else np.nan)
            group_stds.append(np.std(group, ddof=1) if group.size > 1 else np.nan)
        ax.errorbar(range(len(TCEP_COLUMNS)), group_means, yerr=group_stds,
                    color="0.3", lw=0.7, capsize=2, fmt="_", ms=7, zorder=4)
        ax.set_title(f"pH {ph:g}")
        ax.set_xticks(range(len(TCEP_COLUMNS)))
        ax.set_xticklabels(concentration_labels)
        ax.set_xlim(-0.5, 3.5)
        if axis_index == 0:
            ax.set_ylabel(f"A$_{{{WAVELENGTH_NM}}}$")
        ph_results = [result for result in results if result.ph == ph]
        zprime_legend = "\n".join(
            f"{result.tcep_mM:g} mM: Z' = {result.z_prime:.2f}"
            for result in ph_results
        )
        ax.text(
            0.98, 0.98, zprime_legend, transform=ax.transAxes,
            ha="right", va="top", fontsize=5,
            bbox={"facecolor": "white", "edgecolor": "none", "alpha": 0.8, "pad": 1.5},
        )

    z_ax = figure.add_subplot(grid[1, :])
    concentration_positions = {1.0: 0, 5.0: 1, 10.0: 2}
    for ph in PH_ROWS.values():
        ph_results = sorted((result for result in results if result.ph == ph), key=lambda result: result.tcep_mM)
        concentrations = [concentration_positions[result.tcep_mM] for result in ph_results]
        z_primes = np.asarray([result.z_prime for result in ph_results], dtype=float)
        z_ax.plot(
            concentrations, z_primes, "o-", color=PH_COLORS[ph], lw=0.7, ms=3,
            label=f"pH {ph:g}",
        )
        label_offset = {6.0: (5, 4), 7.0: (5, -2), 8.0: (5, -5)}[ph]
        z_ax.annotate(
            f"pH {ph:g}", xy=(concentrations[-1], z_primes[-1]),
            xytext=label_offset, textcoords="offset points", va="center",
            fontsize=6, color=PH_COLORS[ph],
        )
    z_ax.axhline(0.5, color="0.3", lw=0.8, ls="--")
    z_ax.text(
        0.02, 0.52, "robust threshold", transform=z_ax.get_yaxis_transform(),
        ha="left", va="bottom", fontsize=5, color="0.3",
    )
    z_ax.set_xticks(list(concentration_positions.values()))
    z_ax.set_xticklabels(["1", "5", "10"])
    z_ax.set_xlabel("TCEP (mM)")
    z_ax.set_ylabel("Z' factor")
    z_ax.set_ylim(0.0, 1.02)
    z_ax.set_yticks([0.0, 0.25, 0.5, 0.75, 1.0])
    z_ax.set_xlim(-0.15, 2.30)

    if title:
        figure.suptitle(title, fontsize=7)
    figure.subplots_adjust(left=0.08, right=0.98, bottom=0.13, top=0.96)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(out_path, format="pdf")
    figure.savefig(out_path.with_suffix(".png"), format="png", dpi=600)
    plt.close(figure)


def write_summary(results: list[ZPrimeResult], out_path: Path) -> None:
    """Write one row per pH/TCEP comparison, including the underlying control values."""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow([
            "pH", "tcep_mM", "n_negative", "negative_mean_A600", "negative_sd_A600",
            "n_positive", "positive_mean_A600", "positive_sd_A600", "z_prime",
            "negative_wells_A600", "positive_wells_A600",
        ])
        for result in results:
            writer.writerow([
                result.ph, result.tcep_mM,
                np.isfinite(result.negative_values).sum(), result.negative_mean, result.negative_std,
                np.isfinite(result.positive_values).sum(), result.positive_mean, result.positive_std,
                result.z_prime,
                "; ".join(f"{value:.4f}" for value in result.negative_values if np.isfinite(value)),
                "; ".join(f"{value:.4f}" for value in result.positive_values if np.isfinite(value)),
            ])


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("xlsx", type=Path, help="BioTek endpoint .xlsx file.")
    parser.add_argument("--sheet", default=None, help="Worksheet name (default: first sheet).")
    parser.add_argument("--output", type=Path, default=None,
                        help="Output PDF path (default: <xlsx_parent>/../outputs/<stem>_tcep_zprime.pdf).")
    parser.add_argument("--title", default="", help="Optional figure title.")
    args = parser.parse_args()

    plate = load_endpoint_plate(args.xlsx, args.sheet)
    results = analyze_tcep_z_prime(plate)
    out_path = args.output or (args.xlsx.resolve().parent.parent / DEFAULT_OUTPUT_SUBDIR /
                               f"{args.xlsx.stem}_tcep_zprime.pdf")
    plot_tcep_z_prime(plate, results, out_path, title=args.title)
    write_summary(results, out_path.with_suffix(".csv"))

    print("pH  TCEP (mM)  Z' factor")
    for result in results:
        print(f"{result.ph:>2.0f}  {result.tcep_mM:>9g}  {result.z_prime:>9.3f}")
    print(f"\nWrote figure: {out_path}")
    print(f"Wrote table:  {out_path.with_suffix('.csv')}")


if __name__ == "__main__":
    main()
