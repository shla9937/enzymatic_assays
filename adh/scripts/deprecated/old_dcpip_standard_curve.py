"""Fit a DCPIP absorbance standard curve from a BioTek Synergy 384-well plate export.

Layout assumed (per user protocol):
    * Rows A, B, C = triplicate serial dilutions.
    * Columns 1-12 = 1:1 dilution series starting at `start_conc_uM` (default 500 uM DCPIP).
    * Column 12 of each replicate row is the blank (used per-row for baseline correction).
    * Columns 13-24 are empty and ignored.
    * Non-numeric readings (e.g. "OVRFLW") are dropped automatically.

Beer-Lambert is fit as A_corrected = eps_app * [DCPIP], forced through origin,
so `eps_app` (units: AU / uM) can be reused later to convert plate absorbance
into DCPIP concentration for kcat calculations.

Usage:
    python dcpip_standard_curve.py ../test/260731_dcpip_standard_curve.xlsx
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from scipy import stats


REPLICATE_ROWS = ("A", "B", "C")
N_TITRATION_WELLS = 12
BLANK_COLUMN = 12
DEFAULT_START_CONC_UM = 500.0
DEFAULT_OUTPUT_SUBDIR = "outputs"


@dataclass
class StandardCurveFit:
    """Result of a DCPIP standard-curve fit."""

    concentrations_uM: np.ndarray          # per-well concentrations included in fit
    absorbance_corrected: np.ndarray       # per-well baseline-subtracted absorbance
    replicate_labels: np.ndarray           # e.g. "A", "B", "C" per point
    eps_app: float                         # slope, AU / uM (path-length folded in)
    eps_app_stderr: float
    r_squared: float
    n_points: int
    blanks: dict[str, float]               # per-row blank absorbance actually used

    def absorbance_to_concentration(self, absorbance: np.ndarray | float,
                                    blank: float = 0.0) -> np.ndarray | float:
        """Convert raw absorbance to DCPIP concentration (uM) using the fit."""
        return (np.asarray(absorbance) - blank) / self.eps_app


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("xlsx", type=Path, help="Path to the plate reader .xlsx export.")
    parser.add_argument("-o", "--output", type=Path, default=None,
                        help=f"Output PDF path (default: <xlsx_parent>/../{DEFAULT_OUTPUT_SUBDIR}/<stem>_standard_curve.pdf).")
    parser.add_argument("--start-conc", type=float, default=DEFAULT_START_CONC_UM,
                        help=f"Starting DCPIP concentration in µM (default {DEFAULT_START_CONC_UM}).")
    parser.add_argument("--sheet", type=str, default=None,
                        help="Sheet name (default: first sheet).")
    parser.add_argument("--title", type=str, default="",
                        help="Optional plot title.")
    args = parser.parse_args()

    fit = fit_standard_curve(args.xlsx, start_conc_uM=args.start_conc, sheet=args.sheet)

    out_svg = args.output
    if out_svg is None:
        # Sibling `outputs/` dir next to the input's parent (e.g. test/foo.xlsx -> outputs/foo_standard_curve.pdf).
        out_dir = args.xlsx.resolve().parent.parent / DEFAULT_OUTPUT_SUBDIR
        out_svg = out_dir / f"{args.xlsx.stem}_standard_curve.pdf"

    plot_standard_curve(fit, out_svg, title=args.title or args.xlsx.stem)
    _print_summary(fit)
    print(f"\nWrote editable PDF: {out_svg}")


def _find_plate_grid(ws) -> tuple[int, int]:
    """Return (header_row, first_data_col) for the plate grid.

    The BioTek export puts a header row like (None, None, 1, 2, ..., 24, None)
    somewhere in the sheet; row labels A..P sit one column left of the numbers.
    """
    for row_idx, vals in enumerate(ws.iter_rows(values_only=True), start=1):
        for start in range(len(vals) - 1):
            if vals[start] == 1 and vals[start + 1] == 2:
                # Row labels live one column to the left of the "1" column.
                return row_idx, start
    raise ValueError("Could not find plate column-header row (1, 2, ...) in sheet.")


def _load_plate(xlsx_path: Path, sheet: str | None = None) -> dict[str, list]:
    """Load the plate as a dict mapping row label -> list of 24 raw cell values."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    header_row, first_data_col = _find_plate_grid(ws)

    plate: dict[str, list] = {}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        label = row[first_data_col - 1]
        if not isinstance(label, str) or len(label) != 1 or not label.isalpha():
            continue
        plate[label.upper()] = list(row[first_data_col:first_data_col + 24])
    return plate


def _to_float_or_nan(x) -> float:
    if isinstance(x, (int, float)) and not isinstance(x, bool):
        return float(x)
    return np.nan


def fit_standard_curve(
    xlsx_path: str | Path,
    start_conc_uM: float = DEFAULT_START_CONC_UM,
    replicate_rows: tuple[str, ...] = REPLICATE_ROWS,
    n_titration_wells: int = N_TITRATION_WELLS,
    blank_column: int = BLANK_COLUMN,
    sheet: str | None = None,
) -> StandardCurveFit:
    """Load the plate, baseline-correct, and fit a Beer-Lambert line through origin."""
    plate = _load_plate(Path(xlsx_path), sheet=sheet)

    # 1:1 dilutions across wells 1..n-1; well n is the blank (excluded from x).
    dilution_wells = np.arange(1, n_titration_wells)  # wells 1..11
    concs = start_conc_uM / (2.0 ** (dilution_wells - 1))

    xs: list[float] = []
    ys: list[float] = []
    labels: list[str] = []
    blanks: dict[str, float] = {}

    for row_label in replicate_rows:
        if row_label not in plate:
            raise KeyError(f"Replicate row {row_label!r} not present in plate.")
        row_vals = plate[row_label]
        blank = _to_float_or_nan(row_vals[blank_column - 1])
        if np.isnan(blank):
            raise ValueError(f"Blank well {row_label}{blank_column} is not numeric.")
        blanks[row_label] = blank

        for well_idx, conc in zip(dilution_wells, concs):
            raw = _to_float_or_nan(row_vals[well_idx - 1])
            if np.isnan(raw):
                continue  # skip OVRFLW / empty
            xs.append(float(conc))
            ys.append(raw - blank)
            labels.append(row_label)

    x = np.asarray(xs)
    y = np.asarray(ys)

    # Linear regression forced through origin: y = m*x.
    # slope MLE = sum(xy)/sum(x^2); stderr and R^2 computed accordingly.
    sxx = float(np.sum(x * x))
    sxy = float(np.sum(x * y))
    slope = sxy / sxx
    residuals = y - slope * x
    dof = len(x) - 1
    resid_var = float(np.sum(residuals ** 2)) / dof if dof > 0 else np.nan
    slope_se = float(np.sqrt(resid_var / sxx)) if dof > 0 else np.nan
    ss_tot = float(np.sum((y - np.mean(y)) ** 2))
    ss_res = float(np.sum(residuals ** 2))
    r_squared = 1.0 - ss_res / ss_tot if ss_tot > 0 else np.nan

    return StandardCurveFit(
        concentrations_uM=x,
        absorbance_corrected=y,
        replicate_labels=np.asarray(labels),
        eps_app=slope,
        eps_app_stderr=slope_se,
        r_squared=r_squared,
        n_points=len(x),
        blanks=blanks,
    )


def plot_standard_curve(fit: StandardCurveFit, out_svg: Path, title: str = "") -> None:
    """Render an editable PDF of the standard curve with fit line + stats."""
    # Sized for one panel of a 2x2 Nature figure (double-column ~183 mm -> ~88 mm per panel).
    plt.rcParams.update({
        "svg.fonttype": "none",   # keep text editable in SVG (Illustrator/Inkscape)
        "pdf.fonttype": 42,       # keep text editable in PDF (Illustrator/Inkscape)
        "font.family": "sans-serif",
        "font.sans-serif": ["Helvetica", "Arial", "DejaVu Sans"],
        "font.size": 6,
        "axes.labelsize": 6,
        "axes.titlesize": 6,
        "xtick.labelsize": 6,
        "ytick.labelsize": 6,
        "legend.fontsize": 6,
        "axes.linewidth": 0.5,
        "xtick.major.width": 0.5,
        "ytick.major.width": 0.5,
        "xtick.major.size": 2.0,
        "ytick.major.size": 2.0,
        "axes.spines.top": False,
        "axes.spines.right": False,
    })

    # 88 mm x 66 mm ~ 3.46" x 2.60" -- fits one quadrant of a 2x2 double-column Nature figure.
    fig, ax = plt.subplots(figsize=(3.46, 2.60))

    colors = {"A": "#1f77b4", "B": "#d62728", "C": "#2ca02c"}
    for rep in np.unique(fit.replicate_labels):
        mask = fit.replicate_labels == rep
        ax.scatter(
            fit.concentrations_uM[mask],
            fit.absorbance_corrected[mask],
            s=10,
            facecolors="none",
            edgecolors=colors.get(rep, "black"),
            linewidths=0.6,
            label=f"Row {rep}",
            zorder=3,
        )

    # Fit line spanning 0 -> just past max included concentration.
    x_max = float(np.max(fit.concentrations_uM)) * 1.05
    x_line = np.linspace(0.0, x_max, 200)
    y_line = fit.eps_app * x_line
    ax.plot(x_line, y_line, color="black", lw=0.6, zorder=2,
            label=f"Fit: A = {fit.eps_app:.4g} · [DCPIP]")

    stats_text = (
        f"$\\varepsilon_{{app}}$ = {fit.eps_app:.4g} ± {fit.eps_app_stderr:.2g} AU/µM\n"
        f"$R^2$ = {fit.r_squared:.4f}\n"
        f"n = {fit.n_points} wells\n"
        f"blanks (col {BLANK_COLUMN}): "
        + ", ".join(f"{k}={v:.3f}" for k, v in fit.blanks.items())
    )
    ax.text(
        0.98, 0.02, stats_text,
        transform=ax.transAxes, ha="right", va="bottom",
        bbox=dict(boxstyle="round,pad=0.25", facecolor="white",
                  edgecolor="lightgray", linewidth=0.4),
    )

    ax.set_xlabel("[DCPIP] (µM)")
    ax.set_ylabel("A$_{600}$ (blank-subtracted)")
    ax.set_title(title or "DCPIP standard curve")
    ax.set_xlim(left=0)
    ax.set_ylim(bottom=0)
    ax.legend(loc="upper left", frameon=False)

    fig.tight_layout()
    out_svg.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_svg, format="pdf")
    fig.savefig(out_svg.with_suffix(".png"), format="png", dpi=600)
    plt.close(fig)


def _print_summary(fit: StandardCurveFit) -> None:
    print(f"n points fit:        {fit.n_points}")
    print(f"eps_app (AU/µM):     {fit.eps_app:.6g}  (SE {fit.eps_app_stderr:.3g})")
    print(f"R^2:                 {fit.r_squared:.5f}")
    print(f"blanks used:         {fit.blanks}")
    print("Included points (conc µM -> A_corr):")
    for rep in np.unique(fit.replicate_labels):
        mask = fit.replicate_labels == rep
        pts = sorted(zip(fit.concentrations_uM[mask], fit.absorbance_corrected[mask]),
                     reverse=True)
        print(f"  Row {rep}: " + ", ".join(f"{c:.3g}→{a:.3f}" for c, a in pts))


if __name__ == "__main__":
    main()
